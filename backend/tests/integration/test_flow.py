"""End-to-end flow through the HTTP API: upload → extract → correct → approve → export/webhook."""

import csv
import io
import json

import httpx
import pytest
import respx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.llm.truth import load_truth
from app.services.container import Services
from app.services.webhooks import SIGNATURE_HEADER, verify_signature
from tests.conftest import FIXED_NOW, fixture_path
from tests.integration.helpers import API, extraction, low_confidence_fields, upload

pytestmark = pytest.mark.integration
ERP = "https://erp.example.com/webhooks/invoices"


def test_full_flow_with_csv_and_xlsx_export(client: TestClient) -> None:
    pdf = fixture_path("invoice_01")
    truth = load_truth(pdf.with_name(pdf.stem + ".truth.json"))["fields"]
    doc_id = upload(client, pdf)
    data = extraction(client, doc_id)["data"]
    field = low_confidence_fields(data)[0]
    client.patch(
        f"{API}/documents/{doc_id}/extraction",
        json={"corrections": [{"field": field, "value": truth[field]["value"]}]},
    )
    assert client.post(f"{API}/documents/{doc_id}/approve").status_code == 200
    other = upload(client, fixture_path("invoice_02"))  # not approved → excluded by default

    res = client.get(f"{API}/exports", params={"format": "csv"})
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/csv")
    assert 'filename="invoicelens-20261001-120000.csv"' in res.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(res.text)))
    assert [r["document_id"] for r in rows] == [doc_id]
    assert rows[0]["invoice_number"] == truth["invoice_number"]["value"]
    assert float(rows[0]["total"]) == truth["total"]["value"]
    assert rows[0][field] == str(truth[field]["value"])

    res = client.get(f"{API}/exports", params={"format": "xlsx", "ids": f"{doc_id},{other}"})
    wb = load_workbook(io.BytesIO(res.content))
    assert wb["Invoices"].max_row == 3
    assert wb["Line items"].max_row - 1 == len(truth["line_items"]) + len(
        load_truth(fixture_path("invoice_02").with_name("invoice_02_modern_se.truth.json"))[
            "fields"
        ]["line_items"]
    )

    assert client.get(f"{API}/exports", params={"ids": "missing"}).status_code == 404
    assert client.get(f"{API}/exports", params={"format": "pdf"}).status_code == 422


def test_webhook_endpoint_crud(client: TestClient) -> None:
    res = client.post(
        f"{API}/webhook-endpoints",
        json={"url": ERP, "secret": "supersecret1234", "description": "ERP"},
    )
    assert res.status_code == 201
    ep = res.json()
    assert ep["secret_hint"] == "••••1234"
    assert "secret" not in ep
    assert client.get(f"{API}/webhook-endpoints").json() == [ep]
    assert client.get(f"{API}/webhook-endpoints/{ep['id']}").json() == ep

    patched = client.patch(
        f"{API}/webhook-endpoints/{ep['id']}", json={"active": False, "description": None}
    ).json()
    assert patched["active"] is False
    assert patched["description"] is None
    patched = client.patch(
        f"{API}/webhook-endpoints/{ep['id']}", json={"url": "https://other.example.com/h"}
    ).json()
    assert patched["url"] == "https://other.example.com/h"

    bad = client.post(f"{API}/webhook-endpoints", json={"url": "not a url", "secret": "short"})
    assert bad.status_code == 422
    assert {tuple(e["loc"]) for e in bad.json()["errors"]} == {("body", "url"), ("body", "secret")}

    assert client.delete(f"{API}/webhook-endpoints/{ep['id']}").status_code == 204
    assert client.get(f"{API}/webhook-endpoints/{ep['id']}").status_code == 404
    assert client.get(f"{API}/webhook-endpoints/{ep['id']}/deliveries").status_code == 404


@respx.mock
def test_test_ping(client: TestClient, services: Services) -> None:
    route = respx.post(ERP).mock(return_value=httpx.Response(200))
    ep = client.post(f"{API}/webhook-endpoints", json={"url": ERP, "secret": "pingsecret99"}).json()
    res = client.post(f"{API}/webhook-endpoints/{ep['id']}/test")
    assert res.status_code == 200
    assert res.json()["status"] == "succeeded"
    assert res.json()["event"] == "ping"
    assert (
        json.loads(route.calls.last.request.content)["data"]["message"] == "Hello from InvoiceLens"
    )
    assert client.post(f"{API}/webhook-endpoints/999/test").status_code == 404


@respx.mock
def test_approve_pushes_signed_webhook_with_retry(client: TestClient, services: Services) -> None:
    sleeps: list[float] = []
    services.sleep = sleeps.append
    route = respx.post(ERP).mock(side_effect=[httpx.Response(500), httpx.Response(200)])
    secret = "erp-shared-secret"
    ep = client.post(f"{API}/webhook-endpoints", json={"url": ERP, "secret": secret}).json()
    client.post(f"{API}/webhook-endpoints", json={"url": ERP, "secret": secret, "active": False})

    doc_id = upload(client, fixture_path("invoice_03"))
    client.post(f"{API}/documents/{doc_id}/approve")

    assert route.call_count == 2
    assert sleeps == [0.0]  # backoff base is 0 in tests
    request = route.calls.last.request
    assert verify_signature(
        secret, request.headers[SIGNATURE_HEADER], request.content, FIXED_NOW.timestamp()
    )
    body = json.loads(request.content)
    assert body["event"] == "invoice.approved"
    assert body["data"]["document"]["id"] == doc_id
    assert body["data"]["invoice"]["currency"] == "EUR"

    deliveries = client.get(f"{API}/webhook-endpoints/{ep['id']}/deliveries").json()
    assert len(deliveries) == 1
    assert deliveries[0]["status"] == "succeeded"
    assert deliveries[0]["attempts"] == 2
    assert deliveries[0]["document_id"] == doc_id
