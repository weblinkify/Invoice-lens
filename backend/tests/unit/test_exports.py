import csv
import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.db.models import Document, Extraction
from app.services.exports import INVOICE_COLUMNS, LINE_COLUMNS, safe_cell, to_csv, to_xlsx
from tests.factories import s, valid_invoice


def make_doc(doc_id: str = "doc-1", with_extraction: bool = True) -> Document:
    inv = valid_invoice()
    inv.payment_reference = s('=HYPERLINK("evil")')
    doc = Document(
        id=doc_id,
        filename=f"{doc_id}.pdf",
        content_type="application/pdf",
        size_bytes=1,
        sha256="0" * 64,
        storage_key="k",
        status="approved",
        approved_at=datetime(2026, 10, 1, tzinfo=UTC),
    )
    if with_extraction:
        doc.extraction = Extraction(provider="mock", raw={}, data=inv.model_dump(mode="json"))
    return doc


def test_safe_cell() -> None:
    assert safe_cell("=1+1") == "'=1+1"
    assert safe_cell("@SUM") == "'@SUM"
    assert safe_cell("-12") == "'-12"
    assert safe_cell(-12.0) == -12.0
    assert safe_cell("ACME") == "ACME"
    assert safe_cell(None) is None


def test_csv_export() -> None:
    content = to_csv([make_doc(), make_doc("no-ext", with_extraction=False)]).decode()
    rows = list(csv.DictReader(io.StringIO(content)))
    assert list(rows[0].keys()) == list(INVOICE_COLUMNS)
    assert len(rows) == 1
    row = rows[0]
    assert row["vendor_name"] == "Nordic Widgets Oy"
    assert row["invoice_date"] == "2026-03-01"
    assert row["total"] == "434.0"
    assert row["payment_reference"].startswith("'=")
    assert row["needs_review"] == "False"
    assert row["approved_at"] == "2026-10-01T00:00:00+00:00"


def test_csv_empty_has_header() -> None:
    assert to_csv([]).decode().strip() == ",".join(INVOICE_COLUMNS)


def test_xlsx_export() -> None:
    wb = load_workbook(io.BytesIO(to_xlsx([make_doc()])))
    assert wb.sheetnames == ["Invoices", "Line items"]
    inv_ws, lines_ws = wb["Invoices"], wb["Line items"]
    header = [c.value for c in inv_ws[1]]
    assert header == list(INVOICE_COLUMNS)
    values = dict(zip(header, [c.value for c in inv_ws[2]], strict=True))
    assert values["total"] == 434
    assert values["currency"] == "EUR"
    assert [c.value for c in lines_ws[1]] == list(LINE_COLUMNS)
    assert lines_ws.max_row == 3
    assert [c.value for c in lines_ws[2]][2:5] == [1, "Widget", 2]
    assert inv_ws.freeze_panes == "A2"
