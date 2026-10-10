"""CSV / XLSX export of extracted invoices."""

import csv
import io
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.db.models import Document
from app.domain.schema import LINE_ITEM_FIELDS, InvoiceExtraction

INVOICE_COLUMNS: tuple[str, ...] = (
    "document_id",
    "filename",
    "status",
    "vendor_name",
    "vendor_vat_id",
    "invoice_number",
    "invoice_date",
    "due_date",
    "currency",
    "subtotal",
    "vat_total",
    "total",
    "iban",
    "payment_reference",
    "needs_review",
    "approved_at",
)
LINE_COLUMNS: tuple[str, ...] = ("document_id", "invoice_number", "line", *LINE_ITEM_FIELDS)
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Neutralize spreadsheet formula injection in text cells."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def invoice_rows(docs: Sequence[Document]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    invoices: list[dict[str, Any]] = []
    lines: list[dict[str, Any]] = []
    for doc in docs:
        if doc.extraction is None:
            continue
        inv = InvoiceExtraction.model_validate(doc.extraction.data)
        plain = inv.plain_values()
        invoices.append(
            {
                "document_id": doc.id,
                "filename": doc.filename,
                "status": doc.status,
                **{k: plain[k] for k in INVOICE_COLUMNS if k in plain},
                "needs_review": inv.needs_review,
                "approved_at": doc.approved_at.isoformat() if doc.approved_at else None,
            }
        )
        for n, item in enumerate(plain["line_items"], start=1):
            lines.append(
                {
                    "document_id": doc.id,
                    "invoice_number": plain["invoice_number"],
                    "line": n,
                    **item,
                }
            )
    return invoices, lines


def to_csv(docs: Sequence[Document]) -> bytes:
    invoices, _ = invoice_rows(docs)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=INVOICE_COLUMNS, lineterminator="\n")
    writer.writeheader()
    for row in invoices:
        writer.writerow({k: safe_cell(v) for k, v in row.items()})
    return buf.getvalue().encode("utf-8")


def to_xlsx(docs: Sequence[Document]) -> bytes:
    invoices, lines = invoice_rows(docs)
    wb = Workbook()
    sheets = [("Invoices", INVOICE_COLUMNS, invoices), ("Line items", LINE_COLUMNS, lines)]
    for index, (title, columns, rows) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        assert ws is not None
        ws.title = title
        ws.append(list(columns))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([safe_cell(row.get(c)) for c in columns])
        ws.freeze_panes = "A2"
        for i, column in enumerate(columns, start=1):
            width = max([len(column), *(len(str(r.get(column) or "")) for r in rows)])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
